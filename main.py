import imaplib
import email
import os
import re
from email.header import decode_header
from bs4 import BeautifulSoup
from datetime import datetime
import openpyxl
from openpyxl import load_workbook

# --- Variables d'environnement ---
EMAIL_ACCOUNT = os.getenv("EMAIL_ACCOUNT")
EMAIL_PASSWORD = os.getenv("EMAIL_PASSWORD")
if not EMAIL_ACCOUNT or not EMAIL_PASSWORD:
    raise RuntimeError("Variables d'environnement EMAIL_ACCOUNT ou EMAIL_PASSWORD non définies")

# --- Connexion IMAP ---
mail = imaplib.IMAP4_SSL("imap.gmail.com")
mail.login(EMAIL_ACCOUNT, EMAIL_PASSWORD)

# --- Sélection du label Gmail "banque" ---
status, _ = mail.select("banque")
if status != "OK":
    raise RuntimeError("Impossible de sélectionner le dossier/label 'banque'")

# --- Recherche de tous les mails ---
status, messages = mail.search(None, 'ALL')
if status != "OK":
    raise RuntimeError("Erreur lors de la recherche des mails")
messages = messages[0].split()
print(f"{len(messages)} mails trouvés")

# --- Regex libellé/date/montant ---
pattern = re.compile(
    r'(.+?)\s*(\d{1,2} [a-zéû]+ \d{4})\s*([+-]?\d+[,.]\d{2})\s?€',
    re.IGNORECASE
)

# --- Fonctions utilitaires ---
def get_body(msg):
    body = ""
    for part in msg.walk():
        if part.get_content_maintype() == "multipart":
            continue
        if part.get("Content-Disposition") and "attachment" in part.get("Content-Disposition"):
            continue
        try:
            payload = part.get_payload(decode=True)
            if not payload:
                continue
            charset = part.get_content_charset() or 'utf-8'
            text = payload.decode(charset, errors="ignore")
            if part.get_content_type() == "text/html":
                soup = BeautifulSoup(text, "html.parser")
                text = soup.get_text()
            body += text + "\n"
        except:
            continue
    return body.strip()

def get_subject(msg):
    subject, encoding = decode_header(msg['Subject'])[0]
    if isinstance(subject, bytes):
        subject = subject.decode(encoding or 'utf-8', errors='ignore')
    return subject

# --- Parcours des mails ---
rows_by_year = {}  # dictionnaire : année -> liste de lignes

if messages:
    status, data = mail.fetch(b','.join(messages), '(RFC822)')
    if status != "OK":
        raise RuntimeError("Erreur lors du fetch des mails")

    for i in range(0, len(data), 2):
        if len(data[i]) < 2:
            continue
        msg = email.message_from_bytes(data[i][1])
        subject = get_subject(msg)

        # --- Filtrer uniquement les mails transactionnels ---
        if "Opération liée à l'alerte" not in subject:
            continue

        body = get_body(msg)
        match = pattern.search(body)
        if match:
            libelle = match.group(1).strip()
            date_str = match.group(2)
            montant = float(match.group(3).replace(',', '.'))
            msg_id = msg.get("Message-ID")

            # Extraire l'année de l'opération
            try:
                date_obj = datetime.strptime(date_str, "%d %B %Y")
            except ValueError:
                date_obj = datetime.now()
            year = str(date_obj.year)

            # Ajouter la ligne au dictionnaire
            if year not in rows_by_year:
                rows_by_year[year] = []
            rows_by_year[year].append([date_str, libelle, montant, msg_id])
            print(f"{libelle}, {date_str}, {montant}€ -> feuille {year}")

# --- Mettre à jour Excel ---
excel_file = "/mnt/c/Users/adrag/Desktop/suivi.xlsx"
if os.path.exists(excel_file):
    wb = load_workbook(excel_file)
else:
    wb = openpyxl.Workbook()

start_row = 2
start_col = 31  # Colonne AE

for year, rows in rows_by_year.items():
    # --- Créer la feuille si elle n'existe pas ---
    if year not in wb.sheetnames:
        ws = wb.create_sheet(title=year)
    else:
        ws = wb[year]

    # --- Trouver la première ligne vide à partir de start_row ---
    current_row = start_row
    while ws.cell(row=current_row, column=start_col).value is not None:
        current_row += 1

    # --- Ajouter les lignes ---
    for row in rows:
        for i, value in enumerate(row):
            ws.cell(row=current_row, column=start_col + i, value=value)
        current_row += 1

wb.save(excel_file)
print(f"Excel mis à jour : {excel_file}")

# --- Déconnexion ---
mail.logout()
